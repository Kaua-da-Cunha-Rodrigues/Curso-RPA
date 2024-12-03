from docx import Document
from openpyxl import load_workbook
from docx.shared import Pt
import os

#Lista de alunos
nomeArquivo = "Alunos.xlsx"
#Abre o arquivo da lista de alunos
planilhaDadosAlunos = load_workbook(nomeArquivo)
#Pega a página correspondente
sheetSelecionada = planilhaDadosAlunos['Nomes']

#Pega cada linha, a partir da segunda, da coluna A pegando todos os nomes
for linha in range(2, len(sheetSelecionada["A"]) + 1):
    #Abre o arquivo
    arquivoWord = Document("Certificado1.docx")
    #estilo
    estilo = arquivoWord.styles["Normal"]
    #Pega o nome
    nomeAluno = sheetSelecionada['A%s' % linha].value

    for paragrafo in arquivoWord.paragraphs:

        #Se o nome ainda estiver para formatar, no caso estará como "@nome", ele preenche usando o nome do aluno
        if "@nome" in paragrafo.text:
            paragrafo.text = nomeAluno
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

    caminhoCertificados = "Certificados\\" + nomeAluno + ".docx"
    arquivoWord.save(caminhoCertificados)

print("Aplicão encerrada")