from docx import Document
from openpyxl import load_workbook
from docx.shared import Pt
import os

#Lista de alunos
nomeArquivo = "DadosAlunos.xlsx"
#Abre o arquivo da lista de alunos
planilhaDadosAlunos = load_workbook(nomeArquivo)
#Pega a página correspondente
sheetSelecionada = planilhaDadosAlunos['Nomes']

#Pega cada linha, a partir da segunda, da coluna A pegando todos os nomes
for linha in range(2, len(sheetSelecionada["A"]) + 1):
    #Abre o arquivo
    arquivoWord = Document("Certificado2.docx")
    #estilo
    estilo = arquivoWord.styles["Normal"]
    #Pega o nome
    nomeAluno = sheetSelecionada['A%s' % linha].value
    diaInicio = sheetSelecionada['B%s' % linha].value
    mesInicio = sheetSelecionada["C%s" % linha].value
    ano = sheetSelecionada["D%s" % linha].value
    curso = sheetSelecionada["E%s" % linha].value
    instrutor = sheetSelecionada["F%s" % linha].value

    descricao = f"Concluiu com sucesso o curso de {curso} e Excel como automatizar Processos e Planilhas, com a carga horária de 20 horas, promovido pela escola de Cursos Online em {diaInicio} de  {mesInicio} de {ano}."
    for paragrafo in arquivoWord.paragraphs:
        if "@nome" in paragrafo.text:
            paragrafo.text = nomeAluno
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)


        if "@instrutor" in paragrafo.text:
            paragrafo.text = instrutor
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)


        if "@descricao" in paragrafo.text:
            paragrafo.text = descricao
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)


    caminhoCertificados = "CertificadosDiferentes\\" + nomeAluno + ".docx"
    arquivoWord.save(caminhoCertificados)

print("Aplicão encerrada")